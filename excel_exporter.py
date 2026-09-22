import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

def export_vocab_excel(words, unit_name="Unit 2", filename="Bai_tap_Tu_vung.xlsx"):
    wb = openpyxl.Workbook()
    
    # Sheet 1: BÀI TẬP (Tự Điền)
    ws_exercise = wb.active
    ws_exercise.title = "BÀI TẬP (Tự Điền)"
    
    # Sheet 2: ĐÁP ÁN TRA CỨU
    ws_answers = wb.create_sheet(title="ĐÁP ÁN TRA CỨU")

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    italic_font = Font(name="Calibri", size=10, italic=True)
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Header Sheet 1
    ws_exercise.merge_cells("A1:E1")
    t_cell = ws_exercise["A1"]
    t_cell.value = f"BÀI TẬP ÔN TẬP TỪ VỰNG TIẾNG ANH ({unit_name.upper()})"
    t_cell.font = title_font
    t_cell.fill = header_fill
    t_cell.alignment = center_align
    ws_exercise.row_dimensions[1].height = 35

    ws_exercise.merge_cells("A2:E2")
    g_cell = ws_exercise["A2"]
    g_cell.value = "Hướng dẫn: Nhìn nghĩa tiếng Việt ở Cột C và gõ từ tiếng Anh tương ứng vào Cột D. Cột E sẽ tự động kiểm tra đúng/sai."
    g_cell.font = italic_font
    g_cell.alignment = left_align
    ws_exercise.row_dimensions[2].height = 22

    headers_ex = ["STT", "Loại từ (Gợi ý)", "Nghĩa tiếng Việt", "Nhập từ tiếng Anh vào đây", "Kiểm tra kết quả"]
    ws_exercise.append(headers_ex)
    ws_exercise.row_dimensions[3].height = 25
    for c in range(1, 6):
        cell = ws_exercise.cell(row=3, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Header Sheet 2
    ws_answers.append(["STT", "Loại từ", "Từ tiếng Anh (Đáp án)", "Nghĩa tiếng Việt"])
    ws_answers.row_dimensions[1].height = 25
    for c in range(1, 5):
        cell = ws_answers.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    start_row = 4
    for idx, item in enumerate(words, start=1):
        _, word, word_type, vietnamese, _ = item
        current_row = start_row + idx - 1
        ans_row = idx + 1

        ws_answers.append([idx, word_type, word, vietnamese])
        for c in range(1, 5):
            ws_answers.cell(row=ans_row, column=c).border = border_thin

        formula = f'=IF(D{current_row}="","",IF(TRIM(LOWER(D{current_row}))=TRIM(LOWER(\'ĐÁP ÁN TRA CỨU\'!C{ans_row})),"ĐÚNG ✓","SAI ✗"))'
        ws_exercise.append([idx, word_type, vietnamese, "", formula])

        for c in range(1, 6):
            cell = ws_exercise.cell(row=current_row, column=c)
            cell.border = border_thin
            if c in [1, 2, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(color="006100", bold=True)
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006", bold=True)

    max_row = max(start_row + len(words), 5)
    rule_ok = CellIsRule(operator="equal", formula=['"ĐÚNG ✓"'], stopIfTrue=True, fill=green_fill, font=green_font)
    rule_err = CellIsRule(operator="equal", formula=['"SAI ✗"'], stopIfTrue=True, fill=red_fill, font=red_font)

    ws_exercise.conditional_formatting.add(f"E4:E{max_row}", rule_ok)
    ws_exercise.conditional_formatting.add(f"E4:E{max_row}", rule_err)

    col_widths = {"A": 8, "B": 18, "C": 35, "D": 30, "E": 20}
    for col, width in col_widths.items():
        ws_exercise.column_dimensions[col].width = width

    col_widths_ans = {"A": 8, "B": 18, "C": 30, "D": 35}
    for col, width in col_widths_ans.items():
        ws_answers.column_dimensions[col].width = width

    wb.save(filename)
    return filename
